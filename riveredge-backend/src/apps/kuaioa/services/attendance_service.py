"""月度考勤服务。"""

from __future__ import annotations

import calendar
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from apps.kuaioa.models.attendance import KuaioaAttendanceDay, KuaioaAttendanceSheet
from apps.kuaioa.models.employee import KuaioaEmployeeProfile
from apps.kuaioa.models.leave import KuaioaLeaveRequest
from apps.kuaioa.schemas.attendance import (
    AttendanceBatchMark,
    AttendanceDayUpdate,
    AttendanceSheetCreate,
    AttendanceSheetUpdate,
)
from apps.kuaioa.services.kuaioa_list_core import (
    apply_create_audit_by_user_id,
    build_keyword_q,
    generate_daily_code,
    model_to_dict,
    parse_optional_date,
    touch_updated,
)
from core.utils.timezone_utils import resolve_business_datetime, to_site_date
from infra.exceptions.exceptions import BusinessLogicError, NotFoundError
from infra.models.user import User

_ALLOWED_MARK = frozenset({"normal", "leave", "rest"})
_DEFAULT_HOURS = Decimal("8")


def _parse_year_month(value: str) -> tuple[int, int]:
    text = (value or "").strip()
    if len(text) != 7 or text[4] != "-":
        raise BusinessLogicError("年月格式须为 YYYY-MM")
    year = int(text[:4])
    month = int(text[5:7])
    if month < 1 or month > 12:
        raise BusinessLogicError("月份无效")
    return year, month


def _month_bounds(year: int, month: int) -> tuple[date, date, int]:
    last_day = calendar.monthrange(year, month)[1]
    return date(year, month, 1), date(year, month, last_day), last_day


def _employee_eligible_on(emp: KuaioaEmployeeProfile, work_date: date) -> bool:
    """入职次日起可填；离职次日起不可填。"""
    if emp.hire_date and work_date <= emp.hire_date:
        return False
    if emp.leave_date and work_date > emp.leave_date:
        return False
    return True


def _default_cell_hours(standard: Decimal, mark: str) -> tuple[Decimal, Decimal]:
    if mark in ("leave", "rest"):
        return Decimal("0"), Decimal("0")
    return standard, Decimal("0")


class AttendanceService:
    async def list_sheets(
        self,
        tenant_id: int,
        *,
        keyword: Optional[str] = None,
        year_month: Optional[str] = None,
        workshop_name: Optional[str] = None,
        status: Optional[str] = None,
    ) -> list[dict[str, Any]]:
        q = KuaioaAttendanceSheet.filter(tenant_id=tenant_id, deleted_at__isnull=True)
        if year_month:
            q = q.filter(year_month=year_month.strip())
        if workshop_name:
            q = q.filter(workshop_name=workshop_name.strip())
        if status:
            q = q.filter(status=status)
        if keyword:
            q = q.filter(
                build_keyword_q(keyword, "sheet_code", "workshop_name", "production_line_name")
            )
        rows = await q.order_by("-year_month", "-updated_at")
        return [model_to_dict(row) for row in rows]

    async def get_sheet(self, tenant_id: int, sheet_id: int) -> dict[str, Any]:
        sheet = await self._get_sheet_row(tenant_id, sheet_id)
        item = model_to_dict(sheet)
        days = await KuaioaAttendanceDay.filter(
            tenant_id=tenant_id, sheet_id=sheet_id, deleted_at__isnull=True
        ).order_by("employee_name", "employee_id", "work_date")
        day_dicts = [model_to_dict(d) for d in days]
        item["days"] = day_dicts
        item["employee_summaries"] = self._summarize_employees(day_dicts, bool(sheet.has_night))
        item["attendance_hints"] = await self._build_attendance_hints(
            tenant_id, sheet, day_dicts
        )
        return item

    async def _build_attendance_hints(
        self,
        tenant_id: int,
        sheet: KuaioaAttendanceSheet,
        days: list[dict[str, Any]],
    ) -> dict[str, Any]:
        year, month = _parse_year_month(sheet.year_month)
        first, last, _ = _month_bounds(year, month)
        employees = await self._list_workshop_employees(
            tenant_id, sheet.workshop_name, sheet.production_line_name
        )
        present_ids = {int(d["employee_id"]) for d in days}
        new_hires: list[dict[str, Any]] = []
        for emp in employees:
            if emp.hire_date and first <= emp.hire_date <= last:
                new_hires.append(
                    {
                        "employee_id": int(emp.id),
                        "employee_name": emp.full_name,
                        "hire_date": emp.hire_date.isoformat(),
                    }
                )
        left_blocked: list[dict[str, Any]] = []
        for emp in employees:
            if emp.leave_date and emp.leave_date < last:
                left_blocked.append(
                    {
                        "employee_id": int(emp.id),
                        "employee_name": emp.full_name,
                        "leave_date": emp.leave_date.isoformat(),
                    }
                )
        missing_roster = [
            {
                "employee_id": int(emp.id),
                "employee_name": emp.full_name,
            }
            for emp in employees
            if int(emp.id) not in present_ids
            and _employee_eligible_on(emp, last)
        ]
        return {
            "new_hires": new_hires,
            "left_blocked": left_blocked,
            "missing_roster": missing_roster,
        }

    async def export_sheet_excel(self, tenant_id: int, sheet_id: int) -> BytesIO:
        sheet = await self._get_sheet_row(tenant_id, sheet_id)
        data = await self.get_sheet(tenant_id, sheet_id)
        days = data.get("days") or []
        year, month = _parse_year_month(sheet.year_month)
        _, _, last_day = _month_bounds(year, month)
        has_night = bool(sheet.has_night)

        by_emp: dict[int, dict[str, Any]] = {}
        day_keys: list[int] = list(range(1, last_day + 1))
        for d in days:
            eid = int(d["employee_id"])
            bucket = by_emp.setdefault(
                eid,
                {
                    "name": d.get("employee_name"),
                    "days": {},
                },
            )
            wd = d.get("work_date")
            if isinstance(wd, str):
                day_num = int(wd[8:10])
            else:
                day_num = wd.day
            bucket["days"][day_num] = d

        wb = Workbook()
        ws = wb.active
        ws.title = "考勤"
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6 + last_day)
        ws.cell(row=1, column=1, value=f"{sheet.year_month} {sheet.workshop_name} 考勤表")
        ws.cell(row=1, column=1).font = Font(bold=True)

        headers = ["序号", "姓名", "行"] + [str(d) for d in day_keys] + ["正班", "加班", "考勤合计", "员工签字"]
        for col, h in enumerate(headers, start=1):
            cell = ws.cell(row=2, column=col, value=h)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")

        from openpyxl.utils import get_column_letter

        row_idx = 3
        summary_col = 3 + last_day
        for seq, (eid, bucket) in enumerate(sorted(by_emp.items(), key=lambda x: str(x[1]["name"])), 1):
            time_row = row_idx
            ws.cell(row=row_idx, column=1, value=seq)
            ws.cell(row=row_idx, column=2, value=bucket["name"])
            ws.cell(row=row_idx, column=3, value="时间")
            for day_num in day_keys:
                cell = bucket["days"].get(day_num)
                if not cell:
                    ws.cell(row=row_idx, column=3 + day_num, value="")
                    continue
                mark = str(cell.get("mark") or "normal")
                if mark == "leave":
                    val = "X"
                elif mark == "rest":
                    val = "√"
                else:
                    reg = Decimal(str(cell.get("regular_hours") or 0))
                    val = str(reg) if reg > 0 else ""
                ws.cell(row=row_idx, column=3 + day_num, value=val)
            row_idx += 1
            ot_row = row_idx
            ws.cell(row=row_idx, column=3, value="加班")
            for day_num in day_keys:
                cell = bucket["days"].get(day_num)
                ot = Decimal(str(cell.get("ot_hours") or 0)) if cell else Decimal("0")
                ws.cell(row=row_idx, column=3 + day_num, value=str(ot) if ot > 0 else "")
            row_idx += 1
            if has_night:
                ws.cell(row=row_idx, column=3, value="夜班")
                for day_num in day_keys:
                    cell = bucket["days"].get(day_num)
                    night = cell and cell.get("is_night")
                    ws.cell(row=row_idx, column=3 + day_num, value="☆" if night else "")
                row_idx += 1
            time_end = 3 + last_day
            ws.cell(
                row=time_row,
                column=summary_col + 1,
                value=f"=SUM({get_column_letter(4)}{time_row}:{get_column_letter(time_end)}{time_row})",
            )
            ws.cell(
                row=time_row,
                column=summary_col + 2,
                value=f"=SUM({get_column_letter(4)}{ot_row}:{get_column_letter(time_end)}{ot_row})",
            )
            ws.cell(
                row=time_row,
                column=summary_col + 3,
                value=(
                    f"={get_column_letter(summary_col + 1)}{time_row}"
                    f"+{get_column_letter(summary_col + 2)}{time_row}"
                ),
            )
            ws.cell(row=time_row, column=summary_col + 4, value="")

        stream = BytesIO()
        wb.save(stream)
        stream.seek(0)
        return stream

    def _summarize_employees(
        self, days: list[dict[str, Any]], has_night: bool
    ) -> list[dict[str, Any]]:
        by_emp: dict[int, dict[str, Any]] = {}
        for d in days:
            eid = int(d["employee_id"])
            bucket = by_emp.get(eid)
            if not bucket:
                bucket = {
                    "employee_id": eid,
                    "employee_code": d.get("employee_code"),
                    "employee_name": d.get("employee_name"),
                    "regular_total": Decimal("0"),
                    "ot_total": Decimal("0"),
                    "night_count": 0,
                    "leave_days": 0,
                    "rest_days": 0,
                }
                by_emp[eid] = bucket
            mark = str(d.get("mark") or "normal")
            if mark == "leave":
                bucket["leave_days"] += 1
            elif mark == "rest":
                bucket["rest_days"] += 1
            else:
                bucket["regular_total"] += Decimal(str(d.get("regular_hours") or 0))
                bucket["ot_total"] += Decimal(str(d.get("ot_hours") or 0))
            if has_night and d.get("is_night"):
                bucket["night_count"] += 1
        return list(by_emp.values())

    async def create_sheet(
        self, tenant_id: int, data: AttendanceSheetCreate, user_id: int
    ) -> dict[str, Any]:
        year, month = _parse_year_month(data.year_month)
        workshop = (data.workshop_name or "").strip()
        if not workshop:
            raise BusinessLogicError("车间不能为空")
        line = (data.production_line_name or "").strip() or None
        standard = data.standard_hours if data.standard_hours is not None else _DEFAULT_HOURS
        if standard <= 0:
            raise BusinessLogicError("标准日工时须大于 0")

        exists = await KuaioaAttendanceSheet.filter(
            tenant_id=tenant_id,
            deleted_at__isnull=True,
            year_month=f"{year:04d}-{month:02d}",
            workshop_name=workshop,
            production_line_name=line,
        ).exists()
        if exists:
            raise BusinessLogicError("该车间本月考勤单已存在")

        sheet_code = await generate_daily_code(
            KuaioaAttendanceSheet, tenant_id, "ATT", code_field="sheet_code"
        )
        payload: dict[str, Any] = {
            "tenant_id": tenant_id,
            "sheet_code": sheet_code,
            "year_month": f"{year:04d}-{month:02d}",
            "workshop_name": workshop,
            "production_line_name": line,
            "has_night": bool(data.has_night),
            "standard_hours": standard,
            "status": "draft",
            "notes": data.notes,
        }
        await apply_create_audit_by_user_id(payload, user_id)
        sheet = await KuaioaAttendanceSheet.create(**payload)
        await self._generate_days_for_sheet(sheet)
        return await self.get_sheet(tenant_id, int(sheet.id))

    async def update_sheet(
        self, tenant_id: int, sheet_id: int, data: AttendanceSheetUpdate, user_id: int
    ) -> dict[str, Any]:
        sheet = await self._get_sheet_row(tenant_id, sheet_id)
        if sheet.status == "submitted":
            raise BusinessLogicError("已提交的考勤单不可修改表头")
        payload = data.model_dump(exclude_unset=True)
        for key, value in payload.items():
            setattr(sheet, key, value)
        await touch_updated(sheet, user_id)
        await sheet.save()
        return await self.get_sheet(tenant_id, sheet_id)

    async def delete_sheet(self, tenant_id: int, sheet_id: int, user_id: int) -> None:
        sheet = await self._get_sheet_row(tenant_id, sheet_id)
        if sheet.status == "submitted":
            raise BusinessLogicError("已提交的考勤单不可删除")
        sheet.deleted_at = resolve_business_datetime()
        await touch_updated(sheet, user_id)
        await sheet.save()
        await KuaioaAttendanceDay.filter(
            tenant_id=tenant_id, sheet_id=sheet_id, deleted_at__isnull=True
        ).update(deleted_at=resolve_business_datetime())

    async def refresh_roster(self, tenant_id: int, sheet_id: int, user_id: int) -> dict[str, Any]:
        sheet = await self._get_sheet_row(tenant_id, sheet_id)
        if sheet.status == "submitted":
            raise BusinessLogicError("已提交的考勤单不可刷新名单")
        added = await self._generate_days_for_sheet(sheet, only_missing=True)
        await touch_updated(sheet, user_id)
        await sheet.save()
        result = await self.get_sheet(tenant_id, sheet_id)
        result["roster_added_cells"] = added
        return result

    async def update_day(
        self,
        tenant_id: int,
        sheet_id: int,
        day_id: int,
        data: AttendanceDayUpdate,
        user_id: int,
    ) -> dict[str, Any]:
        sheet = await self._get_sheet_row(tenant_id, sheet_id)
        if sheet.status == "submitted":
            raise BusinessLogicError("已提交的考勤单不可改日格")
        day = await KuaioaAttendanceDay.get_or_none(
            id=day_id, tenant_id=tenant_id, sheet_id=sheet_id, deleted_at__isnull=True
        )
        if not day:
            raise NotFoundError("考勤日格不存在")
        payload = data.model_dump(exclude_unset=True)
        if "mark" in payload and payload["mark"] is not None:
            mark = str(payload["mark"]).strip()
            if mark not in _ALLOWED_MARK:
                raise BusinessLogicError("日格标记无效")
            payload["mark"] = mark
            if mark in ("leave", "rest"):
                payload.setdefault("regular_hours", Decimal("0"))
                payload.setdefault("ot_hours", Decimal("0"))
        if "is_night" in payload and payload["is_night"] and not sheet.has_night:
            raise BusinessLogicError("当前考勤单未启用夜班模板")
        for key, value in payload.items():
            setattr(day, key, value)
        await touch_updated(day, user_id)
        await day.save()
        await touch_updated(sheet, user_id)
        await sheet.save()
        return model_to_dict(day)

    async def batch_mark(
        self, tenant_id: int, sheet_id: int, data: AttendanceBatchMark, user_id: int
    ) -> dict[str, Any]:
        sheet = await self._get_sheet_row(tenant_id, sheet_id)
        if sheet.status == "submitted":
            raise BusinessLogicError("已提交的考勤单不可改日格")
        work_date = parse_optional_date(data.work_date)
        if not work_date:
            raise BusinessLogicError("日期无效")
        if data.mark and data.mark not in _ALLOWED_MARK:
            raise BusinessLogicError("日格标记无效")
        if data.is_night and not sheet.has_night:
            raise BusinessLogicError("当前考勤单未启用夜班模板")

        q = KuaioaAttendanceDay.filter(
            tenant_id=tenant_id,
            sheet_id=sheet_id,
            work_date=work_date,
            deleted_at__isnull=True,
        )
        if data.employee_ids:
            q = q.filter(employee_id__in=data.employee_ids)
        rows = await q
        standard = Decimal(str(sheet.standard_hours or _DEFAULT_HOURS))
        for day in rows:
            if data.mark is not None:
                day.mark = data.mark
                reg, ot = _default_cell_hours(standard, data.mark)
                day.regular_hours = reg
                day.ot_hours = ot
            if data.is_night is not None:
                day.is_night = bool(data.is_night)
            await touch_updated(day, user_id)
            await day.save()
        await touch_updated(sheet, user_id)
        await sheet.save()
        return await self.get_sheet(tenant_id, sheet_id)

    async def submit_sheet(self, tenant_id: int, sheet_id: int, user: User) -> dict[str, Any]:
        sheet = await self._get_sheet_row(tenant_id, sheet_id)
        if sheet.status == "submitted":
            raise BusinessLogicError("考勤单已提交")
        sheet.status = "submitted"
        sheet.submitted_at = resolve_business_datetime()
        sheet.submitted_by = user.id
        sheet.submitted_by_name = user.full_name or user.username
        await touch_updated(sheet, user.id)
        await sheet.save()
        return await self.get_sheet(tenant_id, sheet_id)

    async def reopen_sheet(self, tenant_id: int, sheet_id: int, user_id: int) -> dict[str, Any]:
        sheet = await self._get_sheet_row(tenant_id, sheet_id)
        if sheet.status != "submitted":
            raise BusinessLogicError("仅已提交考勤单可重新打开")
        sheet.status = "draft"
        sheet.submitted_at = None
        sheet.submitted_by = None
        sheet.submitted_by_name = None
        await touch_updated(sheet, user_id)
        await sheet.save()
        return await self.get_sheet(tenant_id, sheet_id)

    async def apply_leave_to_attendance(
        self, tenant_id: int, leave: KuaioaLeaveRequest, user_id: int
    ) -> int:
        """请假审批通过后回写未提交考勤日格。返回更新格数。

        满勤日（≥标准工时）标记 leave（X）并清零工时；不足一天则保留 normal，
        正常工时 = 标准 − 请假小时。扣款写入首个受影响日格。
        """
        from apps.kuaioa.services.leave_service import leave_hours_for_date

        if not leave.start_at or not leave.end_at:
            return 0
        start = to_site_date(leave.start_at)
        end = to_site_date(leave.end_at)
        if end < start:
            start, end = end, start

        emp = None
        if leave.applicant_id:
            emp = await KuaioaEmployeeProfile.get_or_none(
                tenant_id=tenant_id,
                user_id=leave.applicant_id,
                deleted_at__isnull=True,
            )
        if not emp and leave.applicant_name:
            emp = await KuaioaEmployeeProfile.filter(
                tenant_id=tenant_id,
                full_name=leave.applicant_name.strip(),
                deleted_at__isnull=True,
            ).first()
        if not emp:
            return 0

        deduct_amount = None
        if getattr(leave, "deduct_enabled", False) and leave.deduct_amount is not None:
            deduct_amount = Decimal(str(leave.deduct_amount))
            if deduct_amount < 0:
                deduct_amount = None

        updated = 0
        deduct_applied = False
        cur = start
        while cur <= end:
            days = await KuaioaAttendanceDay.filter(
                tenant_id=tenant_id,
                employee_id=int(emp.id),
                work_date=cur,
                deleted_at__isnull=True,
            )
            for day in days:
                sheet = await KuaioaAttendanceSheet.get_or_none(
                    id=day.sheet_id, tenant_id=tenant_id, deleted_at__isnull=True
                )
                if not sheet or sheet.status == "submitted":
                    continue
                standard = Decimal(str(sheet.standard_hours or _DEFAULT_HOURS))
                hours = leave_hours_for_date(leave, cur, standard)
                if hours <= 0:
                    continue
                if hours >= standard:
                    day.mark = "leave"
                    day.regular_hours = Decimal("0")
                    day.ot_hours = Decimal("0")
                else:
                    day.mark = "normal"
                    day.regular_hours = max(standard - hours, Decimal("0"))
                    day.ot_hours = Decimal("0")
                day.leave_request_id = int(leave.id)
                if deduct_amount is not None and not deduct_applied:
                    day.leave_deduct_amount = deduct_amount
                    deduct_applied = True
                await touch_updated(day, user_id)
                await day.save()
                updated += 1
            cur += timedelta(days=1)
        return updated

    async def _generate_days_for_sheet(
        self, sheet: KuaioaAttendanceSheet, *, only_missing: bool = False
    ) -> int:
        year, month = _parse_year_month(sheet.year_month)
        first, last, _ = _month_bounds(year, month)
        employees = await self._list_workshop_employees(
            int(sheet.tenant_id),
            sheet.workshop_name,
            sheet.production_line_name,
        )
        standard = Decimal(str(sheet.standard_hours or _DEFAULT_HOURS))
        existing: set[tuple[int, date]] = set()
        if only_missing:
            rows = await KuaioaAttendanceDay.filter(
                tenant_id=sheet.tenant_id, sheet_id=sheet.id, deleted_at__isnull=True
            ).only("employee_id", "work_date")
            existing = {(int(r.employee_id), r.work_date) for r in rows}

        created = 0
        for emp in employees:
            cur = first
            while cur <= last:
                if not _employee_eligible_on(emp, cur):
                    cur += timedelta(days=1)
                    continue
                key = (int(emp.id), cur)
                if key in existing:
                    cur += timedelta(days=1)
                    continue
                await KuaioaAttendanceDay.create(
                    tenant_id=sheet.tenant_id,
                    sheet_id=int(sheet.id),
                    employee_id=int(emp.id),
                    employee_code=emp.employee_code,
                    employee_name=emp.full_name,
                    work_date=cur,
                    regular_hours=standard,
                    ot_hours=Decimal("0"),
                    mark="normal",
                    is_night=False,
                )
                created += 1
                cur += timedelta(days=1)
        return created

    async def _list_workshop_employees(
        self, tenant_id: int, workshop_name: str, production_line_name: Optional[str]
    ) -> list[KuaioaEmployeeProfile]:
        q = KuaioaEmployeeProfile.filter(
            tenant_id=tenant_id,
            deleted_at__isnull=True,
            workshop_name=workshop_name,
        )
        if production_line_name:
            q = q.filter(production_line_name=production_line_name)
        return await q.order_by("full_name", "id")

    async def _get_sheet_row(self, tenant_id: int, sheet_id: int) -> KuaioaAttendanceSheet:
        row = await KuaioaAttendanceSheet.get_or_none(
            id=sheet_id, tenant_id=tenant_id, deleted_at__isnull=True
        )
        if not row:
            raise NotFoundError("考勤单不存在")
        return row
